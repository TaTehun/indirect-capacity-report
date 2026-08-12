"""
processors_example.py

Reference implementation of all four processors.
Split into separate modules for production use.

Sections:
  1. air_processor    (processors/air_processor.py)
  2. otr_processor    (processors/otr_processor.py)
  3. vessel_processor (processors/vessel_processor.py)
  4. report_builder   (processors/report_builder.py)
"""

# ============================================================
# 1. AIR PROCESSOR  (processors/air_processor.py)
# ============================================================
import pandas as pd
from datetime import datetime, timedelta, date
from pathlib import Path
import shutil
import openpyxl

# POD code → site label
AIR_POD_MAP = {
    'DFW': 'TX',
    'JFK': 'NJ',
    'LAX': 'CA',
    'YYZ': 'ON',
    'CHI': 'IL',
}
AIR_SITES = ['DFW', 'JFK', 'LAX', 'YYZ', 'CHI']


def air_load_raw(file_path: Path) -> pd.DataFrame:
    df = pd.read_excel(file_path, sheet_name='RawData', header=0, engine='openpyxl')
    df = df[df['ETA WH'].notna()].copy()
    df['ETA WH'] = pd.to_datetime(df['ETA WH'])
    df['_week']  = df['ETA WH'].dt.isocalendar().week.astype(int)
    df['_year']  = df['ETA WH'].dt.isocalendar().year.astype(int)
    return df


def air_agg_week(df: pd.DataFrame, week: int, year: int) -> dict:
    mask = (df['_week'] == week) & (df['_year'] == year)
    sub  = df[mask]
    result = {}
    for pod in AIR_SITES:
        pod_df = sub[sub['POD'] == pod]
        cw     = pod_df['C/W'].sum()
        hawb   = pod_df['HAWB#'].count()
        result[pod] = (round(cw, 2), int(hawb))
    return result


def air_process(file_path: Path, today: datetime = None) -> dict:
    """
    Returns wk0..wk3 (current week, -1, -2, -3) and avg2/avg3/avg4.
    Each week entry: {pod: (cw, hawb_count)}
    Each avg entry:  {pod: (avg_cw, avg_hawb)}
    """
    if today is None:
        today = datetime.today()

    df = air_load_raw(file_path)

    weeks = []
    for i in range(4):
        d = today - timedelta(weeks=i)
        iso = d.isocalendar()
        weeks.append((int(iso.week), int(iso.year)))

    wk_data = {}
    for i, (wk, yr) in enumerate(weeks):
        wk_data[f'wk{i}'] = air_agg_week(df, wk, yr)

    def _avg(pod, n):
        cw_vals   = [wk_data[f'wk{i}'][pod][0] for i in range(n)]
        hawb_vals = [wk_data[f'wk{i}'][pod][1] for i in range(n)]
        return (round(sum(cw_vals) / n, 2), round(sum(hawb_vals) / n, 2))

    for label, n in [('avg2', 2), ('avg3', 3), ('avg4', 4)]:
        wk_data[label] = {pod: _avg(pod, n) for pod in AIR_SITES}

    return wk_data


# ============================================================
# 2. OTR PROCESSOR  (processors/otr_processor.py)
# ============================================================

OTR_SITES = ['CA', 'IL', 'NJ', 'ON', 'TX']


def otr_load_raw(file_path: Path) -> pd.DataFrame:
    df = pd.read_excel(file_path, sheet_name='RawData', header=0, engine='openpyxl')
    df = df.drop_duplicates(subset=['BL#']).copy()
    df = df[df['ETA WH'].notna()].copy()
    df['ETA WH'] = pd.to_datetime(df['ETA WH'])
    df['_week']  = df['ETA WH'].dt.isocalendar().week.astype(int)
    df['_year']  = df['ETA WH'].dt.isocalendar().year.astype(int)
    return df


def otr_agg_week(df: pd.DataFrame, week: int, year: int) -> dict:
    mask = (df['_week'] == week) & (df['_year'] == year)
    sub  = df[mask]
    return {
        state: int(sub[sub['State'] == state]['BL#'].count())
        for state in OTR_SITES
    }


def otr_process(file_path: Path, today: datetime = None) -> dict:
    """
    Returns wk0..wk3 and avg1..avg4.
    avg1 = wk0 (current week as-is)
    avg2..avg4 = rolling averages
    Each week entry: {state: count}
    """
    if today is None:
        today = datetime.today()

    df = otr_load_raw(file_path)

    weeks = []
    for i in range(4):
        d = today - timedelta(weeks=i)
        iso = d.isocalendar()
        weeks.append((int(iso.week), int(iso.year)))

    wk_data = {}
    for i, (wk, yr) in enumerate(weeks):
        wk_data[f'wk{i}'] = otr_agg_week(df, wk, yr)

    def _avg(state, n):
        vals = [wk_data[f'wk{i}'][state] for i in range(n)]
        return round(sum(vals) / n, 2)

    wk_data['avg1'] = {s: wk_data['wk0'][s] for s in OTR_SITES}
    for label, n in [('avg2', 2), ('avg3', 3), ('avg4', 4)]:
        wk_data[label] = {s: _avg(s, n) for s in OTR_SITES}

    return wk_data


# ============================================================
# 3. VESSEL PROCESSOR  (processors/vessel_processor.py)
# ============================================================

VESSEL_SITES = ['CATOR', 'USDAL', 'USLAX', 'USCHI', 'USNYC']


def vessel_load_raw(file_path: Path) -> pd.DataFrame:
    df = pd.read_excel(file_path, sheet_name='RawData', header=0, engine='openpyxl')
    df = df[df['Status'] != 'Delivery Complete'].copy()
    df['WH ETA'] = pd.to_datetime(df['WH ETA'], errors='coerce')
    df = df[df['WH ETA'].notna()].copy()
    df['_week']  = df['WH ETA'].dt.isocalendar().week.astype(int)
    df['_year']  = df['WH ETA'].dt.isocalendar().year.astype(int)
    return df


def vessel_agg_week(df: pd.DataFrame, week: int, year: int) -> dict:
    mask = (df['_week'] == week) & (df['_year'] == year)
    sub  = df[mask]
    result = {}
    for dst in VESSEL_SITES:
        dst_df = sub[sub['Final Destination'] == dst]
        cntr   = int(dst_df['Cntr.#'].count())
        hbl    = int(dst_df['HBL#'].count())
        result[dst] = (cntr, hbl)
    return result


def vessel_process(file_path: Path, today: datetime = None) -> dict:
    """
    Returns wk0..wk3 (current week + next 3 weeks, forward direction).
    Each week entry: {dst: (cntr_count, hbl_count)}
    """
    if today is None:
        today = datetime.today()

    df = vessel_load_raw(file_path)

    weeks = []
    for i in range(4):
        d = today + timedelta(weeks=i)
        iso = d.isocalendar()
        weeks.append((int(iso.week), int(iso.year)))

    wk_data = {}
    for i, (wk, yr) in enumerate(weeks):
        wk_data[f'wk{i}'] = vessel_agg_week(df, wk, yr)

    return wk_data


# ============================================================
# 4. REPORT BUILDER  (processors/report_builder.py)
# ============================================================

# Row map: site_key -> (air_row, ocean_row, otr_row)
ROW_MAP = {
    'CA': (4,  5,  6),
    'TX': (7,  8,  9),
    'IL': (10, 11, 12),
    'NJ': (13, 14, 15),
    'ON': (16, 17, 18),
}

# Column indices (1-based): D=4 E=5 G=7 H=8 J=10 K=11 M=13 N=14
COL_CW  = [4,  7, 10, 13]
COL_HBL = [5,  8, 11, 14]

SHEET = 'Forecasting'


def _write_pair(ws, row, col_cw, col_hbl, cw_val, hbl_val):
    ws.cell(row=row, column=col_cw).value  = cw_val
    ws.cell(row=row, column=col_hbl).value = hbl_val


def _update_week_headers(ws, base_week: int, base_year: int):
    # Merged ranges C2:E2, F2:H2, I2:K2, L2:N2; write to first cell of each
    for i, col in enumerate([3, 6, 9, 12]):
        d = date.fromisocalendar(base_year, base_week, 1) + timedelta(weeks=i)
        ws.cell(row=2, column=col).value = f'WK{d.isocalendar()[1]}'


def write(
    template_path: Path,
    output_path: Path,
    air_data: dict,
    otr_data: dict,
    vessel_data: dict,
    today: datetime = None,
) -> Path:
    """Copy template → output_path, fill all cells, save. Returns output_path."""
    if today is None:
        today = datetime.today()

    shutil.copy2(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb[SHEET]

    iso       = today.isocalendar()
    base_week = int(iso.week)
    base_year = int(iso.year)

    _update_week_headers(ws, base_week, base_year)

    # AIR: wk0=actual, avg2/avg3/avg4=trend
    air_pod_to_site = {'LAX': 'CA', 'DFW': 'TX', 'CHI': 'IL', 'JFK': 'NJ', 'YYZ': 'ON'}
    for slot_key, col_cw, col_hbl in [
        ('wk0', COL_CW[0], COL_HBL[0]), ('avg2', COL_CW[1], COL_HBL[1]),
        ('avg3', COL_CW[2], COL_HBL[2]), ('avg4', COL_CW[3], COL_HBL[3]),
    ]:
        slot = air_data.get(slot_key, {})
        for pod, site in air_pod_to_site.items():
            cw, hbl = slot.get(pod, (0, 0))
            _write_pair(ws, ROW_MAP[site][0], col_cw, col_hbl, cw, hbl)

    # OTR: avg1=current, avg2/avg3/avg4=trend
    otr_state_to_site = {'CA': 'CA', 'TX': 'TX', 'IL': 'IL', 'NJ': 'NJ', 'ON': 'ON'}
    for slot_key, col_cw, col_hbl in [
        ('avg1', COL_CW[0], COL_HBL[0]), ('avg2', COL_CW[1], COL_HBL[1]),
        ('avg3', COL_CW[2], COL_HBL[2]), ('avg4', COL_CW[3], COL_HBL[3]),
    ]:
        slot = otr_data.get(slot_key, {})
        for state, site in otr_state_to_site.items():
            count = slot.get(state, 0)
            _write_pair(ws, ROW_MAP[site][2], col_cw, col_hbl, count, count)

    # Vessel: wk0=current, wk1/wk2/wk3=next 3 weeks
    vessel_dst_to_site = {'USLAX': 'CA', 'CATOR': 'ON', 'USDAL': 'TX', 'USCHI': 'IL', 'USNYC': 'NJ'}
    for slot_key, col_cw, col_hbl in [
        ('wk0', COL_CW[0], COL_HBL[0]), ('wk1', COL_CW[1], COL_HBL[1]),
        ('wk2', COL_CW[2], COL_HBL[2]), ('wk3', COL_CW[3], COL_HBL[3]),
    ]:
        slot = vessel_data.get(slot_key, {})
        for dst, site in vessel_dst_to_site.items():
            cbm, hbl = slot.get(dst, (0, 0))
            _write_pair(ws, ROW_MAP[site][1], col_cw, col_hbl, cbm, hbl)

    wb.save(output_path)
    return output_path
